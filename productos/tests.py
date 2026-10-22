"""Tests de la sección Productos (Sprint 2). Correr con: python manage.py test"""
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .excel import limpiar_codigo, limpiar_precio
from .models import Categoria, Producto


def crear_usuario(nombre, grupo):
    u = User.objects.create_user(nombre, password="Clave12345!")
    u.groups.add(Group.objects.get(name=grupo))
    return u


def excel_subible(filas, encabezado=("codigo", "nombre", "precio", "categoria")):
    libro = Workbook()
    hoja = libro.active
    hoja.append(list(encabezado))
    for fila in filas:
        hoja.append(list(fila))
    salida = BytesIO()
    libro.save(salida)
    return SimpleUploadedFile("productos.xlsx", salida.getvalue())


class Base(TestCase):
    def setUp(self):
        User.objects.all().delete()  # sin los usuarios iniciales de la migración
        self.admin = crear_usuario("admin_test", "Administrador")
        self.cajero = crear_usuario("cajero_test", "Cajero")
        self.client.force_login(self.admin)


class Permisos(Base):
    def test_cajero_no_entra_a_productos(self):
        self.client.force_login(self.cajero)
        for nombre in ("productos:lista", "productos:nuevo", "productos:importar", "productos:exportar"):
            self.assertEqual(self.client.get(reverse(nombre)).status_code, 403, nombre)


class CargaManual(Base):
    def test_alta_con_categoria_nueva(self):
        self.client.post(reverse("productos:nuevo"), {
            "codigo": " 7790001 ", "nombre": "Coca 500", "precio": "1500.50", "nueva_categoria": "Bebidas"})
        p = Producto.objects.get(codigo="7790001")
        self.assertEqual(p.precio, Decimal("1500.50"))
        self.assertEqual(p.categoria.nombre, "Bebidas")

    def test_codigo_repetido_no_se_guarda(self):
        Producto.objects.create(codigo="1", nombre="A", precio=10)
        r = self.client.post(reverse("productos:nuevo"), {"codigo": "1", "nombre": "B", "precio": "20"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(Producto.objects.count(), 1)

    def test_precio_cero_no_se_guarda(self):
        self.client.post(reverse("productos:nuevo"), {"codigo": "2", "nombre": "B", "precio": "0"})
        self.assertFalse(Producto.objects.exists())

    def test_editar_y_eliminar(self):
        p = Producto.objects.create(codigo="3", nombre="Alfajor", precio=800)
        self.client.post(reverse("productos:editar", args=[p.pk]), {"codigo": "3", "nombre": "Alfajor", "precio": "900"})
        p.refresh_from_db()
        self.assertEqual(p.precio, Decimal("900"))
        self.client.post(reverse("productos:eliminar", args=[p.pk]))
        self.assertFalse(Producto.objects.exists())

    def test_buscador(self):
        Producto.objects.create(codigo="111", nombre="Coca cola", precio=10)
        Producto.objects.create(codigo="222", nombre="Alfajor", precio=10)
        r = self.client.get(reverse("productos:lista"), {"q": "coca", "parcial": "1"})
        self.assertContains(r, "Coca cola")
        self.assertNotContains(r, "Alfajor")


class LimpiezaDeValores(TestCase):
    def test_codigos_que_manda_excel(self):
        self.assertEqual(limpiar_codigo(7790001234567), "7790001234567")
        self.assertEqual(limpiar_codigo(7790001234567.0), "7790001234567")
        self.assertEqual(limpiar_codigo("7790001234567.0"), "7790001234567")
        self.assertEqual(limpiar_codigo(" 0012345 "), "0012345")

    def test_precios_formato_argentino(self):
        self.assertEqual(limpiar_precio("$ 1.234,50"), Decimal("1234.50"))
        self.assertEqual(limpiar_precio("1234,5"), Decimal("1234.50"))
        self.assertEqual(limpiar_precio(1500), Decimal("1500.00"))
        for malo in ("", "abc", -5, 0):
            with self.assertRaises(ValueError):
                limpiar_precio(malo)


class ImportarExcel(Base):
    def test_crea_actualiza_y_reporta_errores(self):
        Producto.objects.create(codigo="100", nombre="Viejo", precio=10)
        archivo = excel_subible([
            (100, "Actualizado", 20, "Bebidas"),        # existe -> actualiza
            (7790001234567.0, "Nuevo", "1.500,00", ""),  # nuevo, código como número de Excel
            ("", "Sin código", 10, ""),                 # error
            ("300", "Precio malo", "abc", ""),          # error
            ("100", "Repetido", 5, ""),                 # error: repetido en la planilla
        ])
        r = self.client.post(reverse("productos:importar"), {"archivo": archivo})
        resumen = r.context["resumen"]
        self.assertEqual(resumen["creados"], 1)
        self.assertEqual(resumen["actualizados"], 1)
        self.assertEqual([fila for fila, _ in resumen["errores"]], [4, 5, 6])
        viejo = Producto.objects.get(codigo="100")
        self.assertEqual((viejo.nombre, viejo.precio, viejo.categoria.nombre), ("Actualizado", Decimal("20"), "Bebidas"))
        self.assertEqual(Producto.objects.get(codigo="7790001234567").precio, Decimal("1500"))

    def test_no_borra_productos_que_no_estan_en_la_planilla(self):
        Producto.objects.create(codigo="999", nombre="Queda", precio=10)
        self.client.post(reverse("productos:importar"), {"archivo": excel_subible([("1", "Otro", 5, "")])})
        self.assertTrue(Producto.objects.filter(codigo="999").exists())

    def test_faltan_columnas(self):
        archivo = excel_subible([("1", "x")], encabezado=("codigo", "nombre"))
        r = self.client.post(reverse("productos:importar"), {"archivo": archivo})
        self.assertIn("precio", r.context["resumen"]["errores"][0][1])

    def test_rechaza_archivos_que_no_son_xlsx(self):
        r = self.client.post(reverse("productos:importar"),
                             {"archivo": SimpleUploadedFile("lista.csv", b"codigo,nombre,precio")})
        self.assertIsNone(r.context["resumen"])
        self.assertFalse(r.context["form"].is_valid())


class ExportarExcel(Base):
    def test_descarga_y_se_puede_volver_a_subir(self):
        cat = Categoria.objects.create(nombre="Golosinas")
        Producto.objects.create(codigo="0012345", nombre="Chicle", precio=Decimal("150.50"), categoria=cat)
        r = self.client.get(reverse("productos:exportar"))
        self.assertIn("attachment", r["Content-Disposition"])
        hoja = load_workbook(BytesIO(r.content)).active
        self.assertEqual([c.value for c in hoja[1]], ["codigo", "nombre", "precio", "categoria"])
        self.assertEqual([c.value for c in hoja[2]], ["0012345", "Chicle", 150.5, "Golosinas"])

        # Ida y vuelta: se sube el mismo archivo y no cambia nada
        r = self.client.post(reverse("productos:importar"),
                             {"archivo": SimpleUploadedFile("p.xlsx", r.content)})
        self.assertEqual(r.context["resumen"]["actualizados"], 1)
        self.assertEqual(r.context["resumen"]["errores"], [])
        self.assertEqual(Producto.objects.get().codigo, "0012345")

    def test_planilla_modelo(self):
        r = self.client.get(reverse("productos:planilla_modelo"))
        hoja = load_workbook(BytesIO(r.content)).active
        self.assertEqual([c.value for c in hoja[1]], ["codigo", "nombre", "precio", "categoria"])
