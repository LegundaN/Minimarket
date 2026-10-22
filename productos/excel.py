"""
Importar y exportar productos con Excel.

Columnas de la planilla:  codigo | nombre | precio | categoria
- Si el código ya existe, se actualizan nombre, precio y categoría.
- Si no existe, se crea el producto.
- La importación NUNCA borra productos.
- Si viene una categoría que no existe, se crea sola.
"""
import re
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .models import Categoria, Producto

COLUMNAS = ["codigo", "nombre", "precio", "categoria"]
OBLIGATORIAS = ["codigo", "nombre", "precio"]


# ---------- Limpieza de valores que vienen del Excel ----------

def _normalizar_encabezado(valor):
    texto = str(valor or "").strip().lower()
    for con, sin in (("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u")):
        texto = texto.replace(con, sin)
    return texto


def limpiar_codigo(valor):
    """Devuelve el código como texto. Excel suele mandar los códigos de barras como número
    (7790001234567.0 o incluso 7.79E+12): acá se convierten a '7790001234567'."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        raise ValueError("el código no es válido")
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        if not valor.is_integer():
            raise ValueError("el código tiene decimales")
        return str(int(valor))
    texto = str(valor).strip()
    if re.fullmatch(r"\d+\.0+", texto):  # "7790001234567.0"
        texto = texto.split(".")[0]
    return texto


def limpiar_precio(valor):
    """Acepta números de Excel o texto con formato argentino: '$ 1.234,50' -> 1234.50"""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        raise ValueError("falta el precio")
    if isinstance(valor, bool):
        raise ValueError("el precio no es un número")
    if isinstance(valor, (int, float, Decimal)):
        precio = Decimal(str(valor))
    else:
        texto = str(valor).replace("$", "").replace(" ", "").strip()
        if "," in texto and "." in texto:
            if texto.rfind(",") > texto.rfind("."):  # 1.234,50
                texto = texto.replace(".", "").replace(",", ".")
            else:                                     # 1,234.50
                texto = texto.replace(",", "")
        elif "," in texto:                            # 1234,50
            texto = texto.replace(",", ".")
        try:
            precio = Decimal(texto)
        except InvalidOperation:
            raise ValueError(f"el precio '{valor}' no es un número")
    if not precio.is_finite():
        raise ValueError("el precio no es un número")
    precio = precio.quantize(Decimal("0.01"))
    if precio <= 0:
        raise ValueError("el precio tiene que ser mayor a cero")
    if precio >= Decimal("100000000"):
        raise ValueError("el precio es demasiado grande")
    return precio


# ---------- Importar ----------

def importar_productos(archivo):
    """Procesa la planilla y devuelve un resumen:
    {"creados": int, "actualizados": int, "errores": [(fila, mensaje), ...]}"""
    resumen = {"creados": 0, "actualizados": 0, "errores": []}
    try:
        libro = load_workbook(archivo, read_only=True, data_only=True)
    except Exception:
        resumen["errores"].append((None, "No se pudo leer el archivo. ¿Es un Excel .xlsx válido?"))
        return resumen

    hoja = libro.worksheets[0]
    filas = hoja.iter_rows(values_only=True)
    encabezado = next(filas, None)
    if not encabezado:
        resumen["errores"].append((None, "La planilla está vacía."))
        return resumen

    posiciones = {}
    for i, valor in enumerate(encabezado):
        nombre = _normalizar_encabezado(valor)
        if nombre in COLUMNAS and nombre not in posiciones:
            posiciones[nombre] = i
    faltan = [c for c in OBLIGATORIAS if c not in posiciones]
    if faltan:
        resumen["errores"].append((1, "Faltan las columnas: " + ", ".join(faltan)
                                   + ". La primera fila tiene que decir: codigo | nombre | precio | categoria"))
        return resumen

    def celda(fila, columna):
        i = posiciones.get(columna)
        return fila[i] if i is not None and i < len(fila) else None

    vistos = {}
    categorias = {c.nombre.lower(): c for c in Categoria.objects.all()}

    with transaction.atomic():
        for numero, fila in enumerate(filas, start=2):
            if fila is None or all(v is None or str(v).strip() == "" for v in fila):
                continue  # fila vacía

            try:
                codigo = limpiar_codigo(celda(fila, "codigo"))
                if not codigo:
                    raise ValueError("falta el código")
                if len(codigo) > 50:
                    raise ValueError("el código es demasiado largo")
                nombre = str(celda(fila, "nombre") or "").strip()
                if not nombre:
                    raise ValueError("falta el nombre")
                if len(nombre) > 150:
                    raise ValueError("el nombre es demasiado largo")
                precio = limpiar_precio(celda(fila, "precio"))
                if codigo in vistos:
                    raise ValueError(f"el código {codigo} está repetido (ya aparece en la fila {vistos[codigo]})")
            except ValueError as error:
                resumen["errores"].append((numero, str(error).capitalize()))
                continue
            vistos[codigo] = numero

            categoria = None
            nombre_categoria = str(celda(fila, "categoria") or "").strip()[:60]
            if nombre_categoria:
                categoria = categorias.get(nombre_categoria.lower())
                if categoria is None:
                    categoria = Categoria.objects.create(nombre=nombre_categoria)
                    categorias[nombre_categoria.lower()] = categoria

            _, creado = Producto.objects.update_or_create(
                codigo=codigo,
                defaults={"nombre": nombre, "precio": precio, "categoria": categoria},
            )
            resumen["creados" if creado else "actualizados"] += 1

    libro.close()
    return resumen


# ---------- Exportar ----------

def _libro_con_encabezado():
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Productos"
    hoja.append(COLUMNAS)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="064E3B")
    for letra, ancho in zip("ABCD", (18, 40, 12, 20)):
        hoja.column_dimensions[letra].width = ancho
    hoja.freeze_panes = "A2"
    return libro, hoja


def _a_bytes(libro):
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def exportar_productos(productos):
    """Excel con las mismas columnas que la importación, para editarlo y volver a subirlo."""
    libro, hoja = _libro_con_encabezado()
    for p in productos:
        hoja.append([p.codigo, p.nombre, float(p.precio), p.categoria.nombre if p.categoria else ""])
    for fila in hoja.iter_rows(min_row=2):
        fila[0].number_format = "@"      # código como texto: evita el 7.79E+12
        fila[2].number_format = "0.00"
    return _a_bytes(libro)


def planilla_modelo():
    """Planilla vacía con las columnas correctas y un ejemplo."""
    libro, hoja = _libro_con_encabezado()
    hoja.append(["7790001234567", "Ejemplo: Gaseosa cola 500 ml (borrar esta fila)", 1500.00, "Bebidas"])
    for fila in range(2, 1001):
        hoja.cell(row=fila, column=1).number_format = "@"   # columna código preparada como texto
        hoja.cell(row=fila, column=3).number_format = "0.00"
    return _a_bytes(libro)
